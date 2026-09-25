"""Exigences du guide du mandataire : annexe 4, classement, dévolution.

Chaque test nomme l'obligation qu'il protège, pour qu'un futur changement dise
lequel il enfreint.
"""

import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi import HTTPException

import comptes
import completude
import concours
import depot
import devolution
import export_cnccfp
from database import ROLE_MANDATAIRE
from db.models import Document
from db.session import campaign_session
from db import enums
from models import Depense, Recette
from _fixture import fresh_campaign, teardown, run_tests


def _depense(montant, rubrique="A1", jour="2026-05-02", libelle="Tract"):
    return Depense(date=date.fromisoformat(jour), libelle=libelle, fournisseur="F",
                   montant_ttc=montant, tva=0.0, categorie_cnccfp=rubrique, statut="Payé")


def test_le_concours_en_nature_consomme_le_plafond_sans_la_tresorerie():
    """Un local prêté ne coûte rien mais compte dans le plafond légal."""
    cid = fresh_campaign(plafond=10_000.0)
    try:
        avant = comptes.compute_stats(cid)
        concours.create_concours(cid, concours.ConcoursIn(
            origine="parti", nature="Mise à disposition d'un local",
            valeur_estimee=2_000.0, methode_evaluation="Loyer de marché",
            rubrique_imputation="B1"))
        apres = comptes.compute_stats(cid)

        assert apres["total_concours_nature"] == 2_000.0
        assert apres["consommation_plafond"] > avant["consommation_plafond"]
        assert apres["solde_tresorerie"] == avant["solde_tresorerie"], "aucun décaissement"
        assert apres["reste_a_depenser"] == avant["reste_a_depenser"] - 2_000.0
    finally:
        teardown(cid)


def test_l_annexe_4_1_distingue_les_trois_origines():
    """La commission veut savoir qui a fourni : candidat, parti ou tiers."""
    cid = fresh_campaign()
    try:
        for origine, valeur in (("candidat", 100.0), ("parti", 200.0), ("tiers_pp", 50.0)):
            concours.create_concours(cid, concours.ConcoursIn(
                origine=origine, nature=f"Concours {origine}", valeur_estimee=valeur))

        synthese = concours.synthese(cid)
        assert synthese["total"] == 350.0
        assert synthese["par_origine"]["parti"]["montant"] == 200.0
        assert synthese["par_origine"]["tiers_pp"]["libelle"] == "Tiers personne physique"
    finally:
        teardown(cid)


def test_les_pieces_sortent_dans_l_ordre_de_la_nomenclature():
    """Le guide impose la répartition verticale puis horizontale.

    Un dossier rendu dans l'ordre d'ajout est à reclasser à la main.
    """
    cid = fresh_campaign()
    try:
        # Saisies à rebours de l'ordre attendu, pour que le tri soit visible.
        for rubrique, libelle, prise in (("B2", "Buffet", "mandataire"),
                                         ("A1", "Tract", "parti"),
                                         ("A1", "Affiche", "mandataire")):
            d = _depense(100.0, rubrique, libelle=libelle)
            d.justificatif_path = f"/data/uploads/{libelle}.pdf"
            comptes.create_depense(cid, d, "gconstant", ROLE_MANDATAIRE)
            if prise == "parti":
                with campaign_session(cid) as s:
                    from db.models import Depense as D
                    dep = s.scalars(__import__("sqlalchemy").select(D)
                                    .where(D.nature == libelle)).first()
                    dep.prise_en_charge = enums.PriseEnCharge.parti

        fichiers = [d["fichier"] for d in depot.list_documents(cid)]
        # Mandataire d'abord (A1 puis B2), parti ensuite.
        assert fichiers == ["Affiche.pdf", "Buffet.pdf", "Tract.pdf"]
    finally:
        teardown(cid)


def test_l_export_porte_les_annexes_4_et_la_ventilation_verticale():
    import io
    from openpyxl import load_workbook

    cid = fresh_campaign()
    try:
        comptes.create_depense(cid, _depense(500.0, "A1"), "gconstant", ROLE_MANDATAIRE)
        concours.create_concours(cid, concours.ConcoursIn(
            origine="candidat", nature="Véhicule", valeur_estimee=300.0,
            rubrique_imputation="D1"))

        wb = load_workbook(io.BytesIO(export_cnccfp.generate_xlsx(cid)))
        assert "Concours nature - Annexe 4" in wb.sheetnames
        assert "Concours nature - Annexe 4.1" in wb.sheetnames

        lignes = list(wb["Depenses par prise en charge"].iter_rows(values_only=True))
        rubriques = {l[0]: l for l in lignes[1:]}
        assert rubriques["A1"][1] == 500.0, "payée par le mandataire"
        assert rubriques["D1"][3] == 300.0, "concours en nature"
    finally:
        teardown(cid)


def test_pas_de_devolution_quand_l_excedent_vient_de_l_apport_du_candidat():
    """L'excédent est alors déduit du remboursement, pas dévolu."""
    cid = fresh_campaign()
    try:
        comptes.create_recette(cid, Recette(date=date(2026, 4, 1), nom_donateur="AMARD Gabriel",
                                            adresse="", montant=5_000.0, type="Apport"),
                               "gconstant", ROLE_MANDATAIRE)
        comptes.create_depense(cid, _depense(1_000.0), "gconstant", ROLE_MANDATAIRE)

        e = devolution.etat(cid)
        assert e["excedent"] == 4_000.0
        assert not e["devolution_due"]
        assert "apport personnel" in e["motif"]
    finally:
        teardown(cid)


def test_devolution_due_quand_l_excedent_vient_des_dons():
    cid = fresh_campaign()
    try:
        comptes.create_recette(cid, Recette(date=date(2026, 4, 1), nom_donateur="MARTIN Claire",
                                            adresse="1 rue X", montant=3_000.0, type="Don"),
                               "gconstant", ROLE_MANDATAIRE)
        comptes.create_depense(cid, _depense(1_000.0), "gconstant", ROLE_MANDATAIRE)

        e = devolution.etat(cid)
        assert e["excedent"] == 2_000.0
        assert e["devolution_due"]
        assert e["montant_devolution"] == 2_000.0

        devolution.enregistrer(cid, devolution.DevolutionIn(
            beneficiaire_type="association", beneficiaire_nom="Les Restos du Cœur",
            montant=2_000.0, date_decision="2027-01-15"))
        assert devolution.etat(cid)["devolution"]["beneficiaire_nom"] == "Les Restos du Cœur"
    finally:
        teardown(cid)


def test_nommer_le_beneficiaire_est_exige_sauf_pour_le_fonds():
    cid = fresh_campaign()
    try:
        try:
            devolution.enregistrer(cid, devolution.DevolutionIn(
                beneficiaire_type="association", montant=100.0))
        except HTTPException as e:
            assert e.status_code == 400
        else:
            raise AssertionError("une association doit être nommée")

        # Le fonds pour la vie associative est une destination unique et connue.
        devolution.enregistrer(cid, devolution.DevolutionIn(
            beneficiaire_type="fonds", montant=100.0))
    finally:
        teardown(cid)


def test_le_releve_bancaire_et_les_recepisses_sont_exiges_au_depot():
    """Trois pièces que l'application ne réclamait pas et que le guide impose."""
    cid = fresh_campaign()
    try:
        cles = {s["cle"] for s in completude.evaluer(cid)["sections"]}
        assert "releves" in cles
        assert "pieces_declaratives" in cles

        manquants = " ".join(completude.evaluer(cid)["manquants"])
        assert "relevé bancaire" in manquants
        assert "Récépissé de déclaration de candidature" in manquants
    finally:
        teardown(cid)


if __name__ == "__main__":
    sys.exit(run_tests(globals()))
