"""Main courante (annexe 8) : journal chronologique recettes/dépenses + export.

Tenue obligatoire dès l'ouverture du compte (spec Bloc E). Chaque ligne reprend
les colonnes de l'annexe 8 ; l'export Excel reproduit ces colonnes sur deux
onglets (Recettes / Dépenses).
"""

from __future__ import annotations

import io

from sqlalchemy import select

from db.models import Depense, Donateur, Recette
from db.session import campaign_session, ensure_campaign_db
from db import enums

_MODE_LABEL = {
    enums.ModePaiement.cheque: "Chèque",
    enums.ModePaiement.virement: "Virement",
    enums.ModePaiement.cb: "CB",
    enums.ModePaiement.prelevement: "Prélèvement",
    enums.ModePaiement.especes: "Espèces",
    enums.ModePaiement.plateforme: "Plateforme",
}

_CATEGORIE_LABEL = {
    enums.CategorieRecette.don: "Don",
    enums.CategorieRecette.apport_perso: "Apport personnel",
    enums.CategorieRecette.pret: "Prêt",
    enums.CategorieRecette.contribution_parti: "Contribution parti",
    enums.CategorieRecette.produit_divers: "Produit divers",
    enums.CategorieRecette.collecte: "Collecte",
}


def _mode(m) -> str:
    return _MODE_LABEL.get(m, "")


def _fmt(d) -> str | None:
    if d is None:
        return None
    return d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d)


def journal(campaign_id: str) -> list[dict]:
    """Journal chronologique unifié (recettes + dépenses)."""
    ensure_campaign_db(campaign_id)
    lignes: list[dict] = []
    with campaign_session(campaign_id) as s:
        donateurs = {d.id: d.nom for d in s.scalars(select(Donateur)).all()}
        for r in s.scalars(select(Recette)).all():
            lignes.append({
                "sens": "recette",
                "date": _fmt(r.date_versement),
                "num_piece": r.num_piece,
                "rubrique": r.rubrique_imputation,
                "nature": _CATEGORIE_LABEL.get(r.categorie, ""),
                "tiers": donateurs.get(r.donateur_id),
                "mode": _mode(r.mode),
                "montant": r.montant,
                "num_releve": r.num_releve_bancaire,
            })
        for d in s.scalars(select(Depense)).all():
            lignes.append({
                "sens": "depense",
                "date": _fmt(d.date_reglement),
                "num_piece": d.num_piece,
                "rubrique": d.rubrique_imputation,
                "nature": d.nature,
                "tiers": d.fournisseur,
                "mode": _mode(d.mode),
                "montant": d.montant_ttc,
                "num_releve": d.num_releve_bancaire,
            })
    # Tri chronologique (les dates manquantes en fin).
    lignes.sort(key=lambda x: (x["date"] is None, x["date"] or ""))
    return lignes


def export_annexe8_xlsx(campaign_id: str) -> bytes:
    """Génère l'export Excel de la main courante (annexe 8), 2 onglets."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    ensure_campaign_db(campaign_id)
    wb = Workbook()

    ws_r = wb.active
    ws_r.title = "Recettes"
    ws_r.append(["N° pièce", "Rubrique", "Nature", "Date de versement",
                 "Donateur / origine", "Mode de versement", "Montant (€)", "N° relevé bancaire"])

    ws_d = wb.create_sheet("Dépenses")
    ws_d.append(["N° pièce", "Rubrique", "Nature", "Date de règlement",
                 "Fournisseur", "Mode de règlement", "Montant facture (€)", "N° relevé bancaire"])

    for ws in (ws_r, ws_d):
        for cell in ws[1]:
            cell.font = Font(bold=True)

    for ligne in journal(campaign_id):
        cols = [ligne["num_piece"], ligne["rubrique"], ligne["nature"], ligne["date"],
                ligne["tiers"], ligne["mode"], ligne["montant"], ligne["num_releve"]]
        (ws_r if ligne["sens"] == "recette" else ws_d).append(cols)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
