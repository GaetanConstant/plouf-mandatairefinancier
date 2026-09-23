"""Export au format officiel CNCCFP (« Main courante du mandataire »).

Produit un classeur Excel dont les onglets et colonnes reproduisent la structure
du modèle officiel, rempli depuis les données de l'application. Objectif : que
l'expert-comptable transcrive case par case dans le modèle officiel, sans deviner.

Réutilise les services existants (identite, annexes, maincourante) et lit
directement les recettes par catégorie via l'ORM pour les annexes 7010/7021/7031.
"""

from __future__ import annotations

import io

from sqlalchemy import select

import annexes
import identite
import maincourante
import recus
from db.session import campaign_session, ensure_campaign_db
from db.models import Donateur, Recette, Evenement
from db.helpers import fmt_date, valides as _valides
from db import enums

_MODE_LABEL = {
    enums.ModePaiement.cheque: "CHQ",
    enums.ModePaiement.virement: "VIR",
    enums.ModePaiement.cb: "CB",
    enums.ModePaiement.prelevement: "PRLV",
    enums.ModePaiement.especes: "ESP",
    enums.ModePaiement.plateforme: "Plateforme",
}


def _mode(m) -> str:
    return _MODE_LABEL.get(m, "")


def _recettes_par_categorie(campaign_id: str, categorie) -> list:
    with campaign_session(campaign_id) as s:
        recs = s.scalars(_valides(select(Recette), Recette).where(Recette.categorie == categorie)
                         .order_by(Recette.date_versement)).all()
        donateurs = {d.id: d for d in s.scalars(select(Donateur)).all()}
        return [(r, donateurs.get(r.donateur_id)) for r in recs]


def generate_xlsx(campaign_id: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    ensure_campaign_db(campaign_id)
    ident = identite.get_identite(campaign_id)
    election = ident.get("election") or {}
    candidat = ident.get("candidat") or {}

    # n° de reçu par recette (dons)
    recu_par_recette = {r["recette_id"]: r["numero_formule"]
                        for r in recus.list_recus(campaign_id) if r["statut"] == "delivre"}

    wb = Workbook()
    bold = Font(bold=True)
    titre_font = Font(bold=True, size=13)

    def new_sheet(name):
        ws = wb.create_sheet(name[:31])  # Excel limite à 31 car.
        return ws

    def headers(ws, cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font = bold

    # ── Fiche signalétique ────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Fiche signalétique"
    ws["A1"] = "COMPTE DE CAMPAGNE — Fiche signalétique"
    ws["A1"].font = titre_font
    rows = [
        ("Élection", election.get("libelle")),
        ("Circonscription", election.get("circonscription")),
        ("Nom de la liste", election.get("nom_liste")),
        ("Nuance politique / Parti", election.get("nuance_politique")),
        ("", ""),
        ("Candidat — Civilité", candidat.get("civilite")),
        ("Nom", candidat.get("nom")),
        ("Prénom", candidat.get("prenom")),
        ("Date de naissance", candidat.get("date_naissance")),
        ("Lieu de naissance", candidat.get("lieu_naissance")),
        ("Adresse", candidat.get("adresse_postale")),
        ("Code postal", candidat.get("code_postal")),
        ("Ville", candidat.get("ville")),
        ("Téléphone", candidat.get("tel")),
        ("Email", candidat.get("email")),
        ("Mandat parlementaire", candidat.get("mandat_parlementaire")),
    ]
    for label, val in rows:
        ws.append([label, val])
        ws[ws.max_row][0].font = bold

    # ── Liste et colistiers ───────────────────────────────────────────────
    ws = new_sheet("Liste et colistiers")
    headers(ws, ["Civ.", "Prénom", "Nom", "Mandat Parlementaire", "1er tour", "2ème tour"])
    for c in annexes.list_colistiers(campaign_id):
        ws.append([c["civilite"], c["prenom"], c["nom"], c["mandat_parlementaire"],
                   "Oui" if c["present_tour1"] else "Non",
                   "Oui" if c["present_tour2"] else "Non"])

    # ── Journal du mandataire ─────────────────────────────────────────────
    ws = new_sheet("Journal du mandataire")
    headers(ws, ["Date Mvt", "Réf. Écriture", "n° chèque ou n° de Remise", "Objet",
                 "Dépenses", "Recettes", "Rappro", "Solde", "Imputation comptable",
                 "Type de paiement"])
    for li in maincourante.journal(campaign_id):
        montant = li["montant"] or 0.0
        ws.append([
            li["date"], li["num_piece"], li["num_cheque_remise"], li["nature"],
            montant if li["sens"] == "depense" else None,
            montant if li["sens"] == "recette" else None,
            "Oui" if li["rapprochement"] else "",
            li["solde"], li["rubrique"], li["mode"],
        ])

    # ── 7010 Annexe 1.1 — Dons ────────────────────────────────────────────
    ws = new_sheet("7010 - Dons")
    headers(ws, ["Nom du donateur", "Prénom du donateur", "Nationalité",
                 "Adresse du domicile fiscal", "Code postal", "Ville", "Pays",
                 "Montant du don", "Date du don", "N° de reçu"])
    for r, don in _recettes_par_categorie(campaign_id, enums.CategorieRecette.don):
        ws.append([
            don.nom if don else None, don.prenom if don else None,
            don.nationalite if don else None, don.adresse if don else None,
            don.code_postal if don else None, don.ville if don else None,
            don.pays_residence if don else None,
            r.montant, fmt_date(r.date_versement), recu_par_recette.get(r.id),
        ])

    # ── 7010 Annexe 1.2 — Collectes ───────────────────────────────────────
    ws = new_sheet("7010 - Collectes")
    headers(ws, ["Événement concerné", "Lieu", "Date", "Montant de la collecte",
                 "Date de versement en banque", "Réf écriture"])
    with campaign_session(campaign_id) as s:
        evs = {e.id: e for e in s.scalars(_valides(select(Evenement), Evenement)).all()}
    for r, _ in _recettes_par_categorie(campaign_id, enums.CategorieRecette.collecte):
        ev = evs.get(r.evenement_id)
        ws.append([ev.titre if ev else None, ev.lieu if ev else None,
                   fmt_date(r.date_versement), r.montant, None, r.num_releve_bancaire])

    # ── 7021 Annexe 3.1 — Apports du candidat ─────────────────────────────
    ws = new_sheet("7021 - Apports candidat")
    headers(ws, ["Civ.", "Prénom", "Nom", "Date du versement", "Date de remise en banque",
                 "N° de remise", "N° Relevé bancaire", "Montant", "Type de paiement"])
    for r, don in _recettes_par_categorie(campaign_id, enums.CategorieRecette.apport_perso):
        ws.append([None, don.prenom if don else None, don.nom if don else None,
                   fmt_date(r.date_versement), fmt_date(r.date_remise_banque),
                   r.num_cheque_remise, r.num_releve_bancaire, r.montant, _mode(r.mode)])

    # ── 7031 Annexe 2 — Versements définitifs des partis ──────────────────
    ws = new_sheet("7031 - Versements partis")
    headers(ws, ["N° de compte", "Nom de la formation politique", "Montant", "Date"])
    for r, don in _recettes_par_categorie(campaign_id, enums.CategorieRecette.contribution_parti):
        ws.append(["7031", don.nom if don else None, r.montant, fmt_date(r.date_versement)])

    # ── Emprunts 7022 / 7023 / 7025 ───────────────────────────────────────
    emprunts = annexes.list_emprunts(campaign_id)

    ws = new_sheet("7022 - Emprunts banque")
    headers(ws, ["Nom de l'établissement prêteur", "Pays du siège", "Date du contrat",
                 "Durée (mois)", "Taux d'intérêt (%)", "Montant de l'emprunt"])
    for e in (x for x in emprunts if x["type"] == "banque"):
        ws.append([e["preteur_nom"], e["preteur_pays"], e["date_contrat"],
                   e["duree_mois"], e["taux"], e["montant"]])

    ws = new_sheet("7023 - Emprunts partis")
    headers(ws, ["Nom du parti ou groupement", "Date du contrat", "Durée (mois)",
                 "Taux d'intérêt (%)", "Montant de l'emprunt"])
    for e in (x for x in emprunts if x["type"] == "parti"):
        ws.append([e["preteur_nom"], e["date_contrat"], e["duree_mois"], e["taux"], e["montant"]])

    ws = new_sheet("7025 - Emprunts PP")
    headers(ws, ["Civ.", "Prénom", "Nom", "Pays de résidence", "Date du contrat",
                 "Durée (mois)", "Date de fin", "Taux d'intérêt (%)", "Montant de l'emprunt"])
    for e in (x for x in emprunts if x["type"] == "personne_physique"):
        ws.append([e["preteur_civilite"], e["preteur_prenom"], e["preteur_nom"],
                   e["preteur_pays"], e["date_contrat"], e["duree_mois"],
                   e["date_fin"], e["taux"], e["montant"]])

    # ── Annexe 7 — Équipe de campagne ─────────────────────────────────────
    ws = new_sheet("Equipe - Annexe 7")
    headers(ws, ["Prénom", "Nom", "Fonction"])
    for m in annexes.list_equipe(campaign_id):
        ws.append([m["prenom"], m["nom"], m["fonction"]])

    # Largeur de colonnes auto (cap à 50)
    for sheet in wb.worksheets:
        for col in sheet.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            sheet.column_dimensions[col[0].column_letter].width = min(width + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
