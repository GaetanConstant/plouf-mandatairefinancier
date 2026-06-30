"""
Vérifie la correspondance entre les fichiers de factures
et les pièces enregistrées dans l'onglet 'Municipales - Factures' du fichier Excel.
"""

import logging
from pathlib import Path

import openpyxl

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

FACTURES_DIR = Path(
    "/Users/gaetanscopa/Nextcloud/11 Politique/11 Elections"
    "/2026_Municipales/Villeurbanne Insoumise/2_Mandataire financier/Factures"
)
EXCEL_PATH = Path(__file__).parent / "Budget 2026.xlsx"
SHEET_NAME = "Municipales - Factures"
COL_PIECE = "Numéro de pièce"
COL_STATUS = "Status"


def load_excel_pieces(excel_path: Path, sheet_name: str) -> list[dict]:
    """Charge les lignes de l'onglet Excel et retourne une liste de dicts."""
    wb = openpyxl.load_workbook(excel_path)
    ws = wb[sheet_name]

    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(row):
            rows.append(dict(zip(headers, row)))
    return rows


def collect_fichiers(factures_dir: Path) -> dict[str, Path]:
    """Collecte tous les fichiers dans le dossier Factures (récursif)."""
    fichiers: dict[str, Path] = {}
    for f in factures_dir.rglob("*"):
        if f.is_file():
            fichiers[f.name] = f
    return fichiers


def main() -> None:
    logger.info("Chargement de l'Excel : %s", EXCEL_PATH)
    rows = load_excel_pieces(EXCEL_PATH, SHEET_NAME)

    logger.info("Lecture des fichiers dans : %s", FACTURES_DIR)
    fichiers_disk = collect_fichiers(FACTURES_DIR)

    pieces_excel = {
        row[COL_PIECE]: row
        for row in rows
        if row.get(COL_PIECE)
    }

    # --- Pièces dans l'Excel mais absentes du disque ---
    manquantes_disk = {
        nom: row
        for nom, row in pieces_excel.items()
        if nom not in fichiers_disk
    }

    # --- Fichiers sur le disque mais absents de l'Excel ---
    manquantes_excel = {
        nom: path
        for nom, path in fichiers_disk.items()
        if nom not in pieces_excel
    }

    # --- Correspondances OK ---
    correspondances = {
        nom: (pieces_excel[nom], fichiers_disk[nom])
        for nom in pieces_excel
        if nom in fichiers_disk
    }

    # --- Statuts "ok" (insensible à la casse) ---
    def is_ok(status: str | None) -> bool:
        return (status or "").strip().lower() == "ok"

    # Pièces marquées Ok dans l'Excel mais pas dans le dossier OK/
    pas_dans_ok = {
        nom: (row, fichiers_disk[nom])
        for nom, row in pieces_excel.items()
        if is_ok(row.get(COL_STATUS))
        and nom in fichiers_disk
        and fichiers_disk[nom].parent.name != "OK"
    }
    # Pièces marquées Ok dans l'Excel et absentes du disque
    ok_absentes = {
        nom: row
        for nom, row in manquantes_disk.items()
        if is_ok(row.get(COL_STATUS))
    }

    print("\n" + "=" * 60)
    print(f"BILAN — {SHEET_NAME}")
    print("=" * 60)

    print(f"\n✅  Pièces trouvées sur le disque ({len(correspondances)}) :")
    for nom, (row, path) in sorted(correspondances.items()):
        status = row.get(COL_STATUS, "")
        dossier = path.parent.name
        print(f"   [{status:^6}]  {nom}  →  {dossier}/")

    # --- Alerte principale : Ok dans Excel mais pas dans dossier OK ---
    total_alerte = len(pas_dans_ok) + len(ok_absentes)
    print(f"\n🚨  Marquées 'Ok' dans l'Excel MAIS pas dans le dossier OK/ ({total_alerte}) :")
    if pas_dans_ok:
        print(f"   — Fichier présent mais dans le mauvais dossier ({len(pas_dans_ok)}) :")
        for nom, (row, path) in sorted(pas_dans_ok.items()):
            print(f"      {nom}  →  {path.parent.name}/")
    if ok_absentes:
        print(f"   — Fichier introuvable sur le disque ({len(ok_absentes)}) :")
        for nom, row in sorted(ok_absentes.items()):
            date = row.get("Date", "")
            date_str = date.strftime("%Y-%m-%d") if hasattr(date, "strftime") else str(date or "")
            print(f"      {nom}  ({date_str})")
    if not pas_dans_ok and not ok_absentes:
        print("   Aucune — tout est en ordre !")

    print(f"\n❌  Dans l'Excel MAIS absentes du disque ({len(manquantes_disk)}) :")
    if manquantes_disk:
        for nom, row in sorted(manquantes_disk.items()):
            status = row.get(COL_STATUS, "")
            date = row.get("Date", "")
            date_str = date.strftime("%Y-%m-%d") if hasattr(date, "strftime") else str(date or "")
            print(f"   [{status:^6}]  {nom}  ({date_str})")
    else:
        print("   Aucune")

    print(f"\n⚠️   Sur le disque MAIS absentes de l'Excel ({len(manquantes_excel)}) :")
    if manquantes_excel:
        for nom, path in sorted(manquantes_excel.items()):
            dossier = path.parent.name
            print(f"   {nom}  →  {dossier}/")
    else:
        print("   Aucune")

    print("\n" + "=" * 60)
    print(f"Total Excel : {len(pieces_excel)} | Disque : {len(fichiers_disk)}")
    print("=" * 60 + "\n")


if __name__ == "__main__":
    main()
